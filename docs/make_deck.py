"""해커톤 발표 자료(SCNU-Lens-발표.pptx, 7장: 표지·문제·해결·기술·시연·효과·마무리)를 만든다. 구글 드라이브에 올리면 구글 슬라이드로 열린다."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SHOT = ROOT / "screenshots"
PUB = ROOT.parent / "frontend" / "public"
OUT = ROOT / "SCNU-Lens-발표.pptx"
TMP = ROOT / ".deck-tmp"
TMP.mkdir(exist_ok=True)
SITE = "https://scnuaram.vercel.app"

BLUE = RGBColor(0x00, 0x58, 0xB2)
NAVY = RGBColor(0x0B, 0x25, 0x45)
SKY = RGBColor(0xA9, 0xCB, 0xF2)
SOFT = RGBColor(0xEE, 0xF4, 0xFB)
PAGE = RGBColor(0xF6, 0xF8, 0xFB)
INK = RGBColor(0x1F, 0x23, 0x2B)
MUTED = RGBColor(0x66, 0x6E, 0x7A)
LINE = RGBColor(0xDD, 0xE3, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xD1, 0x43, 0x3A)
REDSOFT = RGBColor(0xFD, 0xEC, 0xEA)
GREEN = RGBColor(0x1E, 0x9E, 0x6A)
GREENSOFT = RGBColor(0xE6, 0xF6, 0xEE)
GRAY = RGBColor(0xB8, 0xC1, 0xCC)
FONT = "Malgun Gothic"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: 문자열, 또는 줄 목록(각 줄은 문자열이나 (문자열, {옵션}) 조각 목록)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [runs]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        pieces = line if isinstance(line, list) else [line]
        for piece in pieces:
            content, opt = (piece, {}) if isinstance(piece, str) else piece
            r = p.add_run()
            r.text = content
            f = r.font
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = opt.get("color", color)
            f.name = FONT
            rpr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = rpr.makeelement(qn(tag), {})
                    rpr.append(el)
                el.set("typeface", FONT)
    return box


def rect(slide, x, y, w, h, color, radius=None, line=None, shadow=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if not shadow:
        shape.shadow.inherit = False
    return shape


def circle(slide, x, y, d, color, label, size=16, fg=WHITE):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = color
    c.line.fill.background()
    c.shadow.inherit = False
    text(slide, x, y, d, d, label, size=size, color=fg, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x, y, w=0.36, h=0.4, color=GRAY):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False


def phone(slide, image, x, y, h):
    """휴대폰 모양: 검은 테두리 안에 캡처(1170x2532)."""
    w = h * 1170 / 2532
    pad = 0.08
    rect(slide, x - pad, y - pad, w + pad * 2, h + pad * 2, RGBColor(0x11, 0x11, 0x14), radius=0.08, shadow=True)
    slide.shapes.add_picture(str(image), Inches(x), Inches(y), Inches(w), Inches(h))
    return w


def browser(slide, image, x, y, w):
    """노트북 화면처럼: 위쪽 주소창 + 화면 캡처."""
    bar = 0.3
    im = Image.open(image)
    h = w * im.height / im.width
    rect(slide, x - 0.06, y - 0.06, w + 0.12, h + bar + 0.12, RGBColor(0x26, 0x2A, 0x33), radius=0.03, shadow=True)
    for i, c in enumerate((RGBColor(0xFF, 0x5F, 0x57), RGBColor(0xFE, 0xBC, 0x2E), RGBColor(0x28, 0xC8, 0x40))):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.12 + i * 0.18), Inches(y + 0.09), Inches(0.11), Inches(0.11))
        dot.fill.solid(); dot.fill.fore_color.rgb = c; dot.line.fill.background(); dot.shadow.inherit = False
    rect(slide, x + 0.75, y + 0.05, w - 1.0, 0.2, RGBColor(0x3A, 0x3F, 0x4A), radius=0.5)
    text(slide, x + 0.9, y + 0.05, w - 1.3, 0.2, "scnuaram.vercel.app", size=9, color=RGBColor(0xC9, 0xCE, 0xD6), anchor=MSO_ANCHOR.MIDDLE)
    slide.shapes.add_picture(str(image), Inches(x), Inches(y + bar), Inches(w), Inches(h))
    return h + bar


def crop(src, box, name):
    path = TMP / name
    Image.open(src).crop(box).save(path)
    return path


def title(slide, head):
    text(slide, 0.65, 0.5, 12.0, 0.8, head, size=30, bold=True)


def page_no(slide, n):
    text(slide, 12.3, 7.0, 0.5, 0.3, str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


# ---------------------------------------------------------------- 1. 표지 (Title)
V2 = SHOT / "v2"
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 0.8, 1.35, 7.5, 0.4, "2026 SCNU OSS·AI 해커톤 · 기초 트랙", size=14, bold=True, color=SKY)
s.shapes.add_picture(str(PUB / "icon-512.png"), Inches(0.8), Inches(2.0), Inches(0.95), Inches(0.95))
text(s, 1.95, 1.98, 6.5, 1.0, "SCNU Lens", size=54, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.82, 3.3, 7.6, 1.4, ["흩어진 학교·교외 공지를 한곳에,", "나에게 맞는 것만 마감 전에"], size=26, bold=True, color=WHITE, spacing=4)
text(s, 0.82, 4.75, 7.6, 0.9, "순천대 공지 + 교외 장학금·자격증·공모전을 한 앱에서 모아 정리하고 알려 주는 캠퍼스 알리미",
     size=15, color=SKY)
text(s, 0.82, 6.2, 7.6, 0.4, [[("이휴단", {"bold": True, "color": WHITE}), ("   ·   scnuaram.vercel.app", {"color": SKY})]], size=14)
phone(s, V2 / "1-for-me-open.png", 9.4, 0.75, 6.0)
notes(s, "안녕하세요, SCNU Lens를 만든 이휴단입니다. SCNU Lens는 여러 곳에 흩어진 순천대 공지와 교외 장학금·자격증·공모전 정보를 "
         "한 앱으로 통합하고, 나에게 맞는 것만 마감 전에 알려 주는 캠퍼스 알리미입니다. (팀 이름이 있으면 이름 옆에 적어 주세요)")

# ---------------------------------------------------------------- 2. 문제 정의 (Problem)
s = prs.slides.add_slide(BLANK)
bg(s, PAGE)
title(s, "정보는 많은데, 여러 곳에 흩어져 있습니다")
stats = [("112곳", "학교 게시판", "본부 2곳 + 학과·부속기관 110곳"),
         ("+4곳", "교외 사이트", "큐넷 · 데이터자격검정 · K-Startup · 한국장학재단"),
         ("55%", "첨부파일 속 정보", "수집한 공지 159건 중 87건은 신청 기간·대상이 HWP·PDF 안에"),
         ("18건", "중복 게시", "같은 공지가 여러 게시판에 따로 올라옴")]
for i, (big, head, body) in enumerate(stats):
    x = 0.65 + (i % 2) * 3.05
    y = 1.6 + (i // 2) * 2.35
    rect(s, x, y, 2.85, 2.1, WHITE, radius=0.07, line=LINE)
    text(s, x + 0.25, y + 0.2, 2.4, 0.7, big, size=30, bold=True, color=BLUE)
    text(s, x + 0.25, y + 0.95, 2.4, 0.4, head, size=14, bold=True)
    text(s, x + 0.25, y + 1.35, 2.45, 0.7, body, size=10.5, color=MUTED)
pains = [
    ("한 번에 모아 볼 곳이 없다", "학교 공지는 게시판마다, 장학·시험·공모전은 사이트마다 따로 있다", WHITE, BLUE),
    ("대상·마감일이 바로 안 보인다", "나에게 해당되는지, 언제까지인지는 긴 본문과 첨부 한글 파일 속에 있다", WHITE, BLUE),
    ("'접수 기간'이라 열렸는지 헷갈린다", "시험·장학은 기간 접수라, 아직 안 열렸는지 이미 끝났는지 한눈에 안 보인다", WHITE, BLUE),
    ("그래서 좋은 기회를 놓친다", "정보가 있어도 제때 못 보면 장학금·공모전·자격증 마감이 지나간다", REDSOFT, RED),
]
for i, (head, body, fill, accent) in enumerate(pains):
    y = 1.6 + i * 1.18
    rect(s, 6.95, y, 5.75, 1.02, fill, radius=0.1, line=None if fill == REDSOFT else LINE)
    circle(s, 7.2, y + 0.26, 0.5, accent, str(i + 1), size=14)
    text(s, 7.9, y + 0.14, 4.7, 0.4, head, size=15.5, bold=True)
    text(s, 7.9, y + 0.52, 4.7, 0.45, body, size=11, color=MUTED)
page_no(s, 2)
notes(s, "학생에게 필요한 정보는 많습니다. 다만 순천대 공지만 해도 본부와 학과·부속기관 게시판 112곳에 나뉘어 있고, "
         "장학금·자격증 시험·공모전은 또 다른 사이트 네 곳에 있어서 한 번에 모아 볼 곳이 없습니다. "
         "실제로 모은 공지 159건 중 87건은 신청 기간이나 대상이 본문이 아니라 첨부 한글 파일 안에 있었고, 18건은 여러 게시판에 중복으로 올라왔습니다. "
         "게다가 시험이나 장학은 '접수 기간'으로 적혀 있어서 지금 신청할 수 있는지 한눈에 안 보입니다. 그래서 정보가 있어도 제때 못 봐서 기회를 놓칩니다.")

# ---------------------------------------------------------------- 3. 해결 방안 (Solution)
s = prs.slides.add_slide(BLANK)
bg(s, PAGE)
title(s, "한 번 설정하면, 찾아가지 않아도 알려 줍니다")
text(s, 0.65, 1.2, 12.0, 0.4, [[("핵심 가치: ", {"bold": True, "color": BLUE}), "여러 곳의 정보를 한 앱으로 통합해 오가는 시간을 줄입니다. 알림으로 받고, 궁금하면 앱에서 자세히 찾아봅니다."]],
     size=14)
cols = [
    (V2 / "1-for-me-open.png", "나에게 맞는 공지", "학과·관심사로 골라 주고\n왜 추천했는지 이유까지"),
    (V2 / "2-upcoming.png", "접수 상태가 한눈에", "'10/12 접수 시작' · '접수중 D-3'\n접수중·접수 예정 필터, 4가지 정렬"),
    (V2 / "3b-detail-scholarship.png", "긴 공지를 한 장으로", "대상·기간·신청방법·지원내용·\n제출서류를 표로 (첨부파일까지 읽음)"),
    (V2 / "4-calendar-open.png", "놓치지 않게 알림", "마감 7·3·1일 전 · 접수 시작 알림\n캘린더에 '지금 접수중' 목록"),
]
ph = 3.75
pw = ph * 1170 / 2532
for i, (img, head, sub) in enumerate(cols):
    x = 0.75 + i * 3.15
    phone(s, img, x + 0.55, 1.85, ph)
    text(s, x - 0.1, 5.8, pw + 1.2, 0.4, head, size=16, bold=True, align=PP_ALIGN.CENTER)
    text(s, x - 0.2, 6.25, pw + 1.4, 0.8, sub, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
page_no(s, 3)
notes(s, "SCNU Lens는 흩어진 정보를 한 앱으로 통합합니다. 처음에 학과와 관심사를 30초 만에 고르면, 그다음부터는 여러 사이트를 오갈 필요가 없습니다. "
         "나에게 맞는 공지를 이유와 함께 보여 주고, 접수 기간인 공지는 '10월 12일 접수 시작', '접수중 D-3'처럼 지금 상태를 표시합니다. "
         "공지를 누르면 첨부파일까지 읽어서 대상·기간·신청방법을 표로 정리해 보여 주고, 마감 7·3·1일 전과 접수 시작 때 휴대폰으로 알려 줍니다.")

# ---------------------------------------------------------------- 4. 핵심 기술 및 아키텍처 (Tech Stack)
s = prs.slides.add_slide(BLANK)
bg(s, PAGE)
title(s, "모으고 → 읽고 → 정리해서 → 전달합니다")
stages = [
    ("수집", "매시간 자동", ["학교 게시판 112곳 크롤링", "(학교 사이트만)", "공공데이터 API 4종", "큐넷·데이터자격·", "K-Startup·한국장학재단"]),
    ("읽기", "본문 + 첨부", ["HWP·HWPX·PDF·DOCX", "첨부파일 텍스트 추출", "같은 공지 중복 묶기"]),
    ("정리", "서버 안 AI + 규칙", ["분야 분류: MiniLM 다국어", "임베딩(서버에서 실행)", "날짜 추출: 접수 시작·마감·", "시험일·발표일", "상세 칸 표: 대상·기간·", "신청방법·문의·주관"]),
    ("전달", "나에게 맞게", ["학과·관심사 점수 + 이유", "웹 푸시(설치 없는 PWA)", "알림함 · 캘린더(.ics)", "밤에는 모아서 아침에"]),
]
for i, (head, tag, lines) in enumerate(stages):
    x = 0.65 + i * 3.12
    rect(s, x, 1.55, 2.8, 3.4, WHITE, radius=0.06, line=LINE)
    circle(s, x + 0.25, 1.8, 0.55, BLUE, str(i + 1), size=15)
    text(s, x + 0.95, 1.8, 1.8, 0.55, head, size=20, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, 2.5, 2.4, 0.35, tag, size=12, bold=True, color=BLUE)
    text(s, x + 0.25, 2.95, 2.5, 2.1, lines, size=11.5, color=INK, spacing=1)
    if i < 3:
        arrow(s, x + 2.83, 3.05, w=0.26)
stack = [("화면", "Next.js · React · TypeScript · PWA → Vercel"),
         ("서버", "Python · FastAPI · SQLite · APScheduler(매시간 수집·1분 알림)"),
         ("운영", "Oracle Cloud 무료 서버(ARM 2코어) · Caddy(https) · systemd"),
         ("품질", "자동 테스트 135개 · 외부 AI API 없음 → AI 비용 0원")]
for i, (label, body) in enumerate(stack):
    x = 0.65 + (i % 2) * 6.2
    y = 5.4 + (i // 2) * 0.72
    rect(s, x, y, 5.95, 0.58, SOFT, radius=0.2)
    text(s, x + 0.25, y, 0.8, 0.58, label, size=12, bold=True, color=BLUE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.0, y, 4.85, 0.58, body, size=11.5, anchor=MSO_ANCHOR.MIDDLE)
page_no(s, 4)
notes(s, "구조는 네 단계입니다. 매시간 학교 게시판 112곳을 크롤링하고, 교외 정보는 공공데이터포털 공식 API 네 종류로 받습니다. "
         "본문뿐 아니라 HWP·PDF 같은 첨부파일 글자까지 읽고, 여러 게시판에 올라온 같은 공지는 하나로 묶습니다. "
         "그다음 서버 안에서 돌아가는 다국어 임베딩 모델로 분야를 나누고, 규칙으로 접수 시작일·마감일·시험일을 뽑고, 대상·기간·신청방법을 표로 만듭니다. "
         "마지막으로 학과·관심사에 맞춰 이유와 함께 보여 주고 웹 푸시로 알립니다. 외부 AI API를 쓰지 않아 AI 비용은 0원이고, 자동 테스트 135개로 규칙을 지킵니다.")

# ---------------------------------------------------------------- 5. 서비스 시연 (Demo)
s = prs.slides.add_slide(BLANK)
bg(s, PAGE)
title(s, "시연: 설정 → 알림 → 교내·교외 공지 확인")
steps = [
    ("학과 설정", "'컴공'만 쳐도 컴퓨터공학전공 · 관심사 고르기"),
    ("알림 받기", "내 학과 공지·마감 알림이 휴대폰에 (녹화 영상)"),
    ("교내 공지", "'나에게' 탭 → 정리 카드 → 학교 원문 열기"),
    ("교외 정보", "장학·자격증: 접수 예정 / 접수중, 접수 시작순"),
    ("캘린더", "날짜를 누르면 그날 마감 · 지금 접수중 목록"),
]
for i, (head, body) in enumerate(steps):
    y = 1.55 + i * 0.86
    circle(s, 0.7, y + 0.08, 0.5, BLUE, str(i + 1), size=14)
    text(s, 1.4, y + 0.02, 4.4, 0.35, head, size=15, bold=True)
    text(s, 1.4, y + 0.38, 4.6, 0.4, body, size=11, color=MUTED)
rect(s, 0.7, 5.95, 5.2, 0.95, WHITE, radius=0.12, line=LINE)
text(s, 0.95, 5.95, 4.8, 0.95, [[("🎬 알림 녹화 영상", {"bold": True}), " — 여기에 영상을 넣어 주세요"],
                                 "나머지는 휴대폰으로 직접 시연"], size=11.5, anchor=MSO_ANCHOR.MIDDLE, color=INK)
phone(s, V2 / "3-detail-period.png", 6.45, 1.5, 5.3)
phone(s, V2 / "2b-scholarship.png", 9.35, 1.5, 5.3)
text(s, 6.2, 6.95, 3.0, 0.35, "시험: 접수 전 · 신청 기간", size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 9.1, 6.95, 3.0, 0.35, "교외 장학금 · 접수중 D-1", size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
btn = rect(s, 11.9, 1.5, 1.2, 0.5, GREEN, radius=0.5)
btn.click_action.hyperlink.address = SITE
text(s, 11.9, 1.5, 1.2, 0.5, "▶ 열기", size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
page_no(s, 5)
notes(s, "지금부터 시연하겠습니다. (1) 학과 입력에 '컴공'만 쳐도 컴퓨터공학전공이 나옵니다. (2) 알림은 녹화해 둔 영상으로 보여 드리겠습니다. "
         "(3) '나에게' 탭에서 내 학과 공지를 누르면 정리 카드가 열리고, 원문 버튼으로 학교 게시판에 바로 갑니다. "
         "(4) 교외 탭에서는 장학금·자격증을 '접수 예정', '접수중'으로 나눠 보고, 접수 시작순으로 정렬할 수 있습니다. "
         "(5) 캘린더에서 날짜를 누르면 그날 마감과 지금 접수중인 목록이 나옵니다. (오른쪽 위 '열기' 버튼 = 사이트)")

# ---------------------------------------------------------------- 6. 기대 효과 및 확장성 (Impact & Future)
s = prs.slides.add_slide(BLANK)
bg(s, PAGE)
title(s, "기대 효과와 앞으로")
text(s, 0.65, 1.35, 5.4, 0.4, "기대 효과", size=18, bold=True, color=BLUE)
impacts = [
    ("확인 시간 단축", "학교 게시판 112곳 + 교외 4곳의 정보를 앱 한 곳에서"),
    ("마감 놓침 감소", "마감 7·3·1일 전 · 접수 시작 알림, 밤에는 모아서 아침에"),
    ("기회 발견", "교내만 보던 학생도 교외 장학·자격증·공모전까지"),
    ("정보 격차 완화", "선배·단톡방 없이도 신입생이 같은 정보를"),
]
for i, (head, body) in enumerate(impacts):
    y = 1.9 + i * 1.18
    rect(s, 0.65, y, 5.4, 1.0, WHITE, radius=0.1, line=LINE)
    text(s, 0.95, y + 0.12, 4.9, 0.4, head, size=15, bold=True)
    text(s, 0.95, y + 0.52, 4.9, 0.4, body, size=11, color=MUTED)
text(s, 6.55, 1.35, 6.2, 0.4, "확장 계획", size=18, bold=True, color=GREEN)
plans = [
    ("대화로 설정하는 LLM", "지금은 선택식. 무료 서버(2코어·12GB)라 대형 LLM은 못 올렸지만, "
                          "LLM을 붙이면 \"장학금 중 성적 기준 없는 것만\"처럼 대화로 규칙을 만들 수 있음"),
    ("첨부 속 표·이미지까지", "HWP·PDF 글자는 이미 읽음. 다음은 표 안의 날짜, 스캔 포스터 이미지(OCR)"),
    ("하루 체크리스트", "시간표·전자출결과 연결해 공강 시간에 할 일 추천 (학교 협조 필요)"),
    ("어학 시험·다른 대학", "토익 등은 공식 API가 없어 '직접 입력 일정'으로 대응 · 게시판 목록만 바꾸면 다른 대학에도"),
]
for i, (head, body) in enumerate(plans):
    y = 1.9 + i * 1.18
    rect(s, 6.55, y, 6.15, 1.0, GREENSOFT, radius=0.1)
    text(s, 6.85, y + 0.1, 5.7, 0.38, head, size=14.5, bold=True)
    text(s, 6.85, y + 0.47, 5.7, 0.5, body, size=10.5, color=MUTED)
text(s, 0.65, 6.72, 12.0, 0.4, [[("비즈니스 모델: ", {"bold": True, "color": INK}),
                                  "운영비 0원 구조(무료 서버·무료 AI 모델) → 학생 무료 · 대학/학과 단위 공식 알리미로 도입"]],
     size=12.5, color=MUTED, align=PP_ALIGN.CENTER)
page_no(s, 6)
notes(s, "기대 효과는 네 가지입니다. 흩어진 정보를 한 곳에서 보니 확인 시간이 줄고, 알림으로 마감을 덜 놓치고, 교외 장학·자격증까지 발견하고, "
         "신입생도 선배 없이 같은 정보를 얻습니다. 앞으로는 첫째, 지금은 무료 서버라 대형 LLM을 못 올렸는데, LLM을 붙이면 대화로 알림 규칙을 만들 수 있습니다. "
         "둘째, 첨부 HWP·PDF 글자는 이미 읽고 있어서, 다음 단계는 표 안의 날짜와 이미지 포스터 인식입니다. "
         "셋째, 시간표·전자출결과 연결해 공강 시간에 할 일을 추천하는 하루 체크리스트로 넓힐 수 있습니다(학교 협조 필요). "
         "넷째, 토익처럼 공식 API가 없는 시험은 직접 입력 일정으로 대응하고, 게시판 목록만 바꾸면 다른 대학에도 쓸 수 있습니다. "
         "운영비가 0원이라 학생에게는 무료로, 대학·학과 단위 공식 알리미로 도입하는 모델을 생각하고 있습니다.")

# ---------------------------------------------------------------- 7. 마무리 (Outro)
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 0.8, 0.9, 6.0, 0.4, "감사합니다", size=16, bold=True, color=SKY)
text(s, 0.8, 1.6, 7.5, 1.9, ["질문 받겠습니다", "휴대폰으로 직접 써 보세요"], size=40, bold=True, color=WHITE, spacing=6)
text(s, 0.82, 3.85, 7.5, 0.9, ["카메라로 오른쪽 QR을 찍으면 바로 열립니다.", "'홈 화면에 추가'하면 앱처럼 알림을 받아요."],
     size=15, color=WHITE, spacing=4)
btn = rect(s, 0.82, 4.95, 3.3, 0.7, GREEN, radius=0.5)
btn.click_action.hyperlink.address = SITE
text(s, 0.82, 4.95, 3.3, 0.7, "▶  지금 열어 보기", size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i, (label, value) in enumerate((("바로 가기", "scnuaram.vercel.app"), ("소스 코드", "github.com/GOBIEBERPA/scnu-lens"),
                                     ("연락처", "이휴단 · scnuhyudan@gmail.com"))):
    text(s, 0.82, 6.0 + i * 0.4, 1.2, 0.3, label, size=12, bold=True, color=SKY)
    text(s, 2.0, 6.0 + i * 0.4, 5.5, 0.3, value, size=12, color=WHITE)
rect(s, 8.9, 1.5, 3.6, 3.6, WHITE, radius=0.06)
qr = s.shapes.add_picture(str(ROOT / "qr-scnuaram.png"), Inches(9.1), Inches(1.7), Inches(3.2), Inches(3.2))
qr.click_action.hyperlink.address = SITE
text(s, 8.9, 5.3, 3.6, 0.4, "scnuaram.vercel.app", size=13, color=SKY, align=PP_ALIGN.CENTER)
notes(s, "들어 주셔서 감사합니다. 화면의 QR로 바로 써 보실 수 있습니다. 질문 받겠습니다.")

# ---------------------------------------------------------------- 3분 발표 대본 (발표자 노트)
# 위의 notes()는 슬라이드별 설명이고, 실제 발표는 이 대본으로 한다(발표 2분 + 시연 1분 남짓).
SCRIPT = [
    "[0:00-0:10 · 10초]\n안녕하세요, SCNU Lens를 만든 이휴단입니다. "
    "SCNU Lens는 흩어진 학교 공지와 교외 장학·자격증 정보를 한곳에 모아, 나에게 맞는 것만 마감 전에 알려 주는 앱입니다.",
    "[0:10-0:35 · 25초]\n학생에게 필요한 정보는 많습니다. 그런데 순천대 공지만 해도 게시판 112곳에 나뉘어 있고, "
    "장학금·자격증 시험·공모전은 또 다른 사이트 네 곳에 있습니다. "
    "실제로 모은 공지의 절반 이상은 신청 기간과 대상이 첨부 한글 파일 안에 있었고, "
    "시험처럼 '접수 기간'으로 적힌 건 지금 신청할 수 있는지 한눈에 안 보입니다. 그래서 정보가 있어도 제때 못 봐서 기회를 놓칩니다.",
    "[0:35-0:55 · 20초]\nSCNU Lens는 이 정보를 한 앱으로 통합합니다. 학과와 관심사를 한 번 고르면 나에게 맞는 공지를 이유와 함께 보여 주고, "
    "'10월 12일 접수 시작', '접수중 D-3'처럼 지금 상태를 표시합니다. 긴 공지는 표 한 장으로 정리하고, 마감 전에 휴대폰으로 알려 줍니다.",
    "[0:55-1:15 · 20초]\n구조는 네 단계입니다. 매시간 학교 게시판과 공공데이터 API에서 모으고, 첨부 HWP·PDF까지 읽고, "
    "서버 안에서 돌아가는 AI 모델과 규칙으로 분류하고 날짜를 뽑아, 웹 푸시로 전달합니다. 외부 AI API를 쓰지 않아 AI 비용은 0원입니다.",
    "[1:15-2:25 · 70초 · 휴대폰 화면으로 전환]\n"
    "① (녹화 영상, 10초) 먼저 알림입니다. 내 학과 공지와 마감 알림이 이렇게 휴대폰에 옵니다.\n"
    "② (마이페이지 → 테스트 알림 보내기) 지금 바로 한 번 보내 보겠습니다. … 이렇게 옵니다. 알림을 누르면 앱이 열립니다.\n"
    "③ ('나에게' 탭 → 공지 하나) 내 학과 공지가 이유와 함께 모여 있고, 누르면 대상·기간·신청방법이 표로 정리돼 있습니다. "
    "버튼 하나로 학교 원문으로 갑니다.\n"
    "④ ('장학·공모·자격증' 탭 → '접수 예정' 칩 → 정렬 '접수 시작순') 교외 장학금과 자격증은 접수중·접수 예정으로 나눠 보고, "
    "곧 열리는 순서로 볼 수 있습니다.\n"
    "⑤ (캘린더 → 오늘) 오늘을 누르면 지금 신청할 수 있는 목록이 마감 가까운 순으로 나옵니다.",
    "[2:25-2:50 · 25초 · 슬라이드로 복귀]\n기대 효과는 흩어진 정보를 한 곳에서 보는 시간 단축, 마감 놓침 감소, 교외 기회 발견입니다. "
    "앞으로는 LLM을 붙여 대화로 알림 규칙을 만들고, 첨부 속 표와 이미지 포스터까지 읽고, "
    "시간표·출결과 연결한 하루 체크리스트로 넓히려 합니다.",
    "[2:50-3:00 · 10초]\n화면의 QR로 지금 바로 써 보실 수 있습니다. 들어 주셔서 감사합니다.",
]
for slide, body in zip(prs.slides, SCRIPT):
    slide.notes_slide.notes_text_frame.text = body

prs.save(OUT)
print("saved", OUT)
